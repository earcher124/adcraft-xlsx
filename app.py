from flask import Flask, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
import io
import re
import base64

app = Flask(__name__)

PURPLE_DARK  = "4C1D95"
PURPLE_MID   = "6B21A8"
PURPLE_LIGHT = "EDE9FE"
PURPLE_PALE  = "FAF8FF"
WHITE        = "FFFFFF"
GREY_LIGHT   = "F3F4F6"
GREY_BORDER  = "DDD6FE"
TEXT_DARK    = "1A1A1A"
TEXT_MID     = "374151"

def side(color=GREY_BORDER, style="thin"):
    return Side(border_style=style, color=color)

def border(color=GREY_BORDER):
    s = side(color)
    return Border(left=s, right=s, top=s, bottom=s)

def parse_plan(text):
    """Parse markdown plan text into structured sections."""
    sections = []
    current_section = None
    current_lines = []

    for raw_line in text.split("\n"):
        line = raw_line.strip()

        if re.match(r'^## ', line):
            if current_section is not None:
                sections.append((current_section, current_lines))
            current_section = line.replace("## ", "").strip()
            current_lines = []
        elif re.match(r'^# ', line):
            if current_section is not None:
                sections.append((current_section, current_lines))
            current_section = line.replace("# ", "").strip()
            current_lines = []
        else:
            current_lines.append(raw_line)

    if current_section is not None:
        sections.append((current_section, current_lines))
    elif current_lines:
        sections.append(("Plan Output", current_lines))

    return sections

def strip_markdown(text):
    """Remove markdown formatting for cell content."""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'^#+\s*', '', text)
    return text.strip()

def build_excel(data):
    wb = Workbook()

    # ── Tab 1: Ad Plan ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Ad Plan"
    ws.sheet_view.showGridLines = False

    # Column widths
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 72

    row = 1

    # Title banner
    ws.merge_cells(f"A{row}:B{row}")
    cell = ws[f"A{row}"]
    cell.value = "AdCraft — Advertising Plan"
    cell.font = Font(name="Arial", size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=PURPLE_DARK)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    ws.row_dimensions[row].height = 36
    row += 1

    # Plan name subtitle
    ws.merge_cells(f"A{row}:B{row}")
    cell = ws[f"A{row}"]
    cell.value = data.get("Plan Name", "")
    cell.font = Font(name="Arial", size=13, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=PURPLE_MID)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    ws.row_dimensions[row].height = 28
    row += 1

    # Spacer
    ws.row_dimensions[row].height = 8
    row += 1

    # Metadata fields
    meta_fields = [
        ("Business Name",    data.get("Business Name", "")),
        ("Primary Goal",     data.get("Primary Goal", "")),
        ("Geography",        data.get("Geography", "")),
        ("Monthly Ad Budget",data.get("Monthly Ad Budget", "")),
        ("Budget Tier",      data.get("Budget Tier", "")),
        ("Status",           data.get("Status", "")),
    ]

    for i, (label, value) in enumerate(meta_fields):
        fill_color = PURPLE_PALE if i % 2 == 0 else WHITE
        # Label cell
        lc = ws[f"A{row}"]
        lc.value = label
        lc.font = Font(name="Arial", size=10, bold=True, color=PURPLE_MID)
        lc.fill = PatternFill("solid", fgColor=fill_color)
        lc.alignment = Alignment(horizontal="left", vertical="center", indent=2, wrap_text=True)
        lc.border = border()
        ws.row_dimensions[row].height = 20

        # Value cell
        vc = ws[f"B{row}"]
        vc.value = str(value) if value else "—"
        vc.font = Font(name="Arial", size=10, color=TEXT_DARK)
        vc.fill = PatternFill("solid", fgColor=fill_color)
        vc.alignment = Alignment(horizontal="left", vertical="center", indent=2, wrap_text=True)
        vc.border = border()
        row += 1

    # Spacer
    ws.row_dimensions[row].height = 12
    row += 1

    # Plan content sections
    plan_text = data.get("Plan Output", "")
    sections = parse_plan(plan_text)

    if not sections:
        sections = [("Plan Output", plan_text.split("\n"))]

    for section_title, lines in sections:
        # Section header row
        ws.merge_cells(f"A{row}:B{row}")
        hc = ws[f"A{row}"]
        hc.value = section_title
        hc.font = Font(name="Arial", size=11, bold=True, color=WHITE)
        hc.fill = PatternFill("solid", fgColor=PURPLE_MID)
        hc.alignment = Alignment(horizontal="left", vertical="center", indent=2)
        ws.row_dimensions[row].height = 24
        row += 1

        # Content rows
        content_buffer = []
        for raw_line in lines:
            line = raw_line.strip()
            if line == "" or re.match(r'^---+$', line):
                if content_buffer:
                    # Write buffered content
                    text = "\n".join(content_buffer)
                    ws.merge_cells(f"A{row}:B{row}")
                    cc = ws[f"A{row}"]
                    cc.value = text
                    cc.font = Font(name="Arial", size=10, color=TEXT_DARK)
                    cc.fill = PatternFill("solid", fgColor=WHITE)
                    cc.alignment = Alignment(horizontal="left", vertical="top", indent=2, wrap_text=True)
                    cc.border = Border(
                        left=side(), right=side(),
                        bottom=side(GREY_BORDER, "hair")
                    )
                    # Estimate row height
                    est_lines = max(len(text.split("\n")), len(text) // 90 + 1)
                    ws.row_dimensions[row].height = max(15, est_lines * 15)
                    row += 1
                    content_buffer = []
                continue

            # Handle ### subheadings inline
            if re.match(r'^### ', line):
                if content_buffer:
                    text = "\n".join(content_buffer)
                    ws.merge_cells(f"A{row}:B{row}")
                    cc = ws[f"A{row}"]
                    cc.value = text
                    cc.font = Font(name="Arial", size=10, color=TEXT_DARK)
                    cc.fill = PatternFill("solid", fgColor=WHITE)
                    cc.alignment = Alignment(horizontal="left", vertical="top", indent=2, wrap_text=True)
                    cc.border = Border(left=side(), right=side(), bottom=side(GREY_BORDER, "hair"))
                    est_lines = max(len(text.split("\n")), len(text) // 90 + 1)
                    ws.row_dimensions[row].height = max(15, est_lines * 15)
                    row += 1
                    content_buffer = []

                subhead = line.replace("### ", "").strip()
                ws.merge_cells(f"A{row}:B{row}")
                sc = ws[f"A{row}"]
                sc.value = subhead
                sc.font = Font(name="Arial", size=10, bold=True, color=PURPLE_DARK)
                sc.fill = PatternFill("solid", fgColor=PURPLE_LIGHT)
                sc.alignment = Alignment(horizontal="left", vertical="center", indent=2)
                ws.row_dimensions[row].height = 20
                row += 1
                continue

            # Handle bullet points
            bullet_match = re.match(r'^[-•]\s+(.+)', line)
            if bullet_match:
                content_buffer.append("• " + strip_markdown(bullet_match.group(1)))
                continue

            numbered_match = re.match(r'^\d+\.\s+(.+)', line)
            if numbered_match:
                content_buffer.append(strip_markdown(line))
                continue

            # Table rows — simplified
            if line.startswith("|") and line.endswith("|"):
                if re.match(r'^\|[-|\s]+\|$', line):
                    continue
                cells = [c.strip() for c in line.split("|") if c.strip()]
                content_buffer.append("  |  ".join(cells))
                continue

            content_buffer.append(strip_markdown(line))

        # Flush remaining buffer
        if content_buffer:
            text = "\n".join(content_buffer)
            ws.merge_cells(f"A{row}:B{row}")
            cc = ws[f"A{row}"]
            cc.value = text
            cc.font = Font(name="Arial", size=10, color=TEXT_DARK)
            cc.fill = PatternFill("solid", fgColor=WHITE)
            cc.alignment = Alignment(horizontal="left", vertical="top", indent=2, wrap_text=True)
            cc.border = Border(left=side(), right=side(), bottom=side(GREY_BORDER, "hair"))
            est_lines = max(len(text.split("\n")), len(text) // 90 + 1)
            ws.row_dimensions[row].height = max(15, est_lines * 15)
            row += 1

        # Section spacer
        ws.row_dimensions[row].height = 10
        row += 1

    # ── Tab 2: Raw Data ─────────────────────────────────────────────
    ws2 = wb.create_sheet("Raw Data")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 100

    # Header
    for col, label in enumerate(["Field", "Value"], 1):
        c = ws2.cell(row=1, column=col, value=label)
        c.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=PURPLE_DARK)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
        ws2.row_dimensions[1].height = 22

    raw_fields = [
        "Plan Name", "Business Name", "Primary Goal",
        "Geography", "Monthly Ad Budget", "Budget Tier", "Status", "Plan Output"
    ]
    for i, field in enumerate(raw_fields):
        r = i + 2
        fill = PURPLE_PALE if i % 2 == 0 else WHITE
        lc = ws2.cell(row=r, column=1, value=field)
        lc.font = Font(name="Arial", size=9, bold=True, color=PURPLE_MID)
        lc.fill = PatternFill("solid", fgColor=fill)
        lc.alignment = Alignment(horizontal="left", vertical="top", indent=2)
        lc.border = border()

        vc = ws2.cell(row=r, column=2, value=str(data.get(field, "")))
        vc.font = Font(name="Arial", size=9, color=TEXT_DARK)
        vc.fill = PatternFill("solid", fgColor=fill)
        vc.alignment = Alignment(horizontal="left", vertical="top", indent=2, wrap_text=True)
        vc.border = border()

        text = str(data.get(field, ""))
        est = max(len(text.split("\n")), len(text) // 110 + 1)
        ws2.row_dimensions[r].height = max(15, est * 14)

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

@app.route("/generate", methods=["POST"])
def generate():
    data = request.get_json(force=True)
    if not data:
        return jsonify({"error": "No data"}), 400

    buf = build_excel(data)
    filename = re.sub(r'[^\w\s-]', '', data.get("Plan Name", "ad-plan")).strip().replace(" ", "-")
    filename = f"{filename}-AdCraft.xlsx"

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "ok"})

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080)
